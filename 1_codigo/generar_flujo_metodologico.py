"""Genera el Capitulo III (Evaluacion de Modelos Base) como .docx, con la
estructura 3.1-3.9, la Figura 3.1 (flujo del capitulo) y la Tabla 3.1
(correspondencia con el metodo de investigacion).

Fuentes de las cifras:
  - Intra-dataset (Mendeley, 10 modelos): 3_resultados_test_por_modelo/
    metricas_pooled_por_modelo.csv y metricas_por_fold.csv (se leen aqui).
  - Hiperparametros: V3_CFG en run_freshness_v3.py.
  - Cross-dataset (Trucha, Rohu; metodo congelado, sin x2): repositorio
    Generalizacion_ViT_DenseNet (3_resultados/, Documentacion_Tecnica_ViT_DenseNet.docx,
    Indicadores_Robustez_TRD_Gap_RBE.docx) -- copiadas en CROSS_DATASET abajo
    porque esos CSV viven en otro repositorio.

Uso (desde la raiz del repo; requiere python-docx y matplotlib):
    python 1_codigo/generar_flujo_metodologico.py
Salida: 6_documentacion/Flujo_Metodologico_Capitulo_III.docx
        5_graficas_informe/flujo_metodologico.png (Figura 3.1 suelta)
"""
import csv
import io
import re
import statistics
import textwrap

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = "6_documentacion/Flujo_Metodologico_Capitulo_III.docx"
FIG_OUT = "5_graficas_informe/flujo_metodologico.png"
CSV_POOLED = "3_resultados_test_por_modelo/metricas_pooled_por_modelo.csv"
CSV_FOLDS = "3_resultados_test_por_modelo/metricas_por_fold.csv"
LOG_ENTRENAMIENTO = "2_logs_entrenamiento/dnc_5fold_resume.log"

NOMBRES = {
    "vit_base_patch16_224": "ViT-Base/16", "densenet121": "DenseNet-121",
    "swin_base_patch4_window7_224": "Swin-Base", "vit_small_patch14_dinov2": "DINOv2-Small",
    "efficientnet_b0": "EfficientNet-B0", "efficientnet_lite0": "EfficientNet-Lite0",
    "mobilenetv3_small_100": "MobileNetV3-Small", "resnet50": "ResNet-50",
    "mobilenetv2_100": "MobileNetV2", "mobilenetv1_100": "MobileNetV1",
}
ELEGIDOS = ("vit_base_patch16_224", "densenet121")

# Cross-dataset, metodo congelado sin x2 (repo Generalizacion_ViT_DenseNet).
# f1: F1 macro oficial (media de 5 folds); f1_pool/acc_pool/auc_pool: pooled de
# la corrida de matrices de confusion/ROC; trd/gap/rbe: Indicadores_Robustez.
CROSS_DATASET = {
    ("vit_base_patch16_224", "Trucha"): dict(f1=86.11, f1_pool=86.18, acc_pool=86.87, auc_pool=0.9197,
                                             trd=88.81, gap=0.1085, rbe=26.59),
    ("vit_base_patch16_224", "Rohu"): dict(f1=88.17, f1_pool=87.67, acc_pool=89.45, auc_pool=0.9514,
                                           trd=90.93, gap=0.0879, rbe=31.38),
    ("densenet121", "Trucha"): dict(f1=80.88, f1_pool=80.80, acc_pool=82.00, auc_pool=0.8973,
                                    trd=84.55, gap=0.1478, rbe=None),
    ("densenet121", "Rohu"): dict(f1=82.85, f1_pool=83.01, acc_pool=85.57, auc_pool=0.9140,
                                  trd=86.61, gap=0.1281, rbe=None),
}
PARAMS_M = {"vit_base_patch16_224": "85.8", "densenet121": "7.0"}


def cargar_intra():
    pooled = {r["model"]: r for r in csv.DictReader(open(CSV_POOLED))}
    f1s = {}
    for r in csv.DictReader(open(CSV_FOLDS)):
        f1s.setdefault(r["model"], []).append(float(r["f1"]))
    for m, r in pooled.items():
        r["f1_media"] = statistics.mean(f1s[m])
        r["f1_de"] = statistics.stdev(f1s[m])
    return sorted(pooled.values(), key=lambda r: -float(r["f1"]))


def epochs_por_modelo():
    """Epochs ejecutados por (modelo, fold) segun el log de entrenamiento. El log es de la
    corrida reanudada: no incluye ViT-Base/16 folds 1-2 ni EfficientNet-Lite0/MobileNetV3-Small."""
    epochs, actual = {}, None
    for linea in open(LOG_ENTRENAMIENTO):
        m = re.match(r"--- (\S+) fold (\d) ---", linea)
        if m:
            actual = (m.group(1), int(m.group(2)))
            epochs[actual] = 0
            continue
        m = re.match(r"\s+epoch\s+(\d+)/\d+\s+val_F1=", linea)
        if m and actual:
            epochs[actual] = int(m.group(1))
    por_modelo = {}
    for (modelo, fold), e in epochs.items():
        por_modelo.setdefault(modelo, {})[fold] = e
    return por_modelo


# --- Figura 3.1 ---------------------------------------------------------------

ETAPAS = [
    ("ETAPA 1", "Preparación", "#2E6DB4", "#EEF4FB", [
        ("3.2", "Conjuntos de datos y preparación de datos",
         "Preparar los tres repositorios con un procesamiento común (depuración, recorte de la región "
         "ocular, homologación de las categorías de frescura y normalización). Mendeley (fuente) se "
         "particiona y solo su conjunto de entrenamiento se aumenta; OneDrive y Kaggle (objetivo) se "
         "reservan hasta la evaluación cross-dataset.",
         "conjuntos de datos comparables entre dominios, sin entrenamiento previo."),
        ("3.3", "Diseño experimental de los modelos candidatos",
         "Definir los modelos candidatos a partir de la literatura, junto con la partición de los datos "
         "(5 folds, semilla 42), los hiperparámetros, las métricas y los escenarios intra-dataset y "
         "cross-dataset (Mendeley → OneDrive y Mendeley → Kaggle).",
         "protocolo experimental común."),
    ]),
    ("ETAPA 2", "Implementación\ny evaluación\nintra-dataset", "#B87A22", "#FDF6EC", [
        ("3.4", "Implementación de los modelos candidatos",
         "Entrenar los modelos de Machine Learning y Deep Learning candidatos bajo condiciones homogéneas "
         "en el dominio fuente (Mendeley).",
         "modelos candidatos entrenados."),
        ("3.5", "Evaluación intra-dataset de los modelos candidatos",
         "Medir el desempeño de cada modelo en el dominio fuente mediante exactitud, precisión, "
         "sensibilidad, F1-score y especificidad.",
         "desempeño comparativo de los modelos candidatos."),
    ]),
    ("ETAPA 3", "Selección", "#2E8A5A", "#EDF6F0", [
        ("3.6", "Selección de los modelos de referencia",
         "Elegir, según el desempeño intra-dataset, los modelos que conformarán la línea base experimental.",
         "modelos de referencia (línea base experimental)."),
    ]),
    ("ETAPA 4", "Evaluación\ncross-dataset\ny análisis", "#7E57D2", "#F4F0FB", [
        ("3.7", "Evaluación cross-dataset de los modelos de referencia",
         "Evaluar los modelos de referencia, entrenados con Mendeley, sobre los dominios objetivo OneDrive "
         "y Kaggle mediante el método congelado (solo se reentrena la cabeza de clasificación), con 5 "
         "folds y semilla 42.",
         "desempeño frente a dominios no vistos durante el entrenamiento base."),
        ("3.8", "Análisis del efecto de la variabilidad visual entre conjuntos de datos",
         "Cuantificar, para cada modelo de referencia, la degradación y la tasa de retención del "
         "desempeño (F1-Score Macro) entre el escenario intra-dataset y los cross-dataset, y comparar "
         "los modelos según su robustez, considerando que el cambio de dominio incluye también "
         "diferencias de especie.",
         "magnitud del efecto del cambio de dominio y robustez comparada de los modelos."),
    ]),
    ("ETAPA 5", "Consolidación", "#2A7B7B", "#ECF5F5", [
        ("3.9", "Consolidación de la línea base experimental",
         "Reunir en una síntesis final los modelos de referencia con sus resultados intra-dataset, "
         "cross-dataset y de retención, que servirán como punto de comparación para el modelo propuesto.",
         "línea base experimental consolidada y diagnóstico del problema."),
    ]),
]
AZUL = "#1F3A5C"
WRAP = 80          # caracteres por linea dentro de cada tarjeta
LH = 0.30          # alto de linea de texto


def figura_flujo() -> io.BytesIO:
    def lineas(t):
        return textwrap.wrap(t, WRAP)

    def alto_tarjeta(s):
        _, _, busca, res = s
        return 0.95 + LH * (len(lineas(busca)) + 1 + len(lineas(res)))

    gap_t, pad, flecha = 0.18, 0.18, 0.55
    alto_etapas = [sum(alto_tarjeta(s) for s in e[4]) + gap_t * (len(e[4]) - 1) + 2 * pad for e in ETAPAS]
    h_cab, h_pie = 1.7, 1.05
    total = h_cab + sum(alto_etapas) + h_pie + flecha * (len(ETAPAS) + 1) + 0.2
    W = 14.0
    fig, ax = plt.subplots(figsize=(W * 0.72, total * 0.72))
    ax.set_xlim(0, W)
    ax.set_ylim(0, total)
    ax.axis("off")

    def caja(x, y, w, h, fc, ec, lw=1.6, r=0.12):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle=f"round,pad=0,rounding_size={r}",
                                    linewidth=lw, edgecolor=ec, facecolor=fc))

    def baja(y):
        ax.annotate("", xy=(W / 2, y - flecha + 0.02), xytext=(W / 2, y),
                    arrowprops=dict(arrowstyle="-|>", color=AZUL, lw=2))
        return y - flecha

    y = total - 0.1
    caja(0.2, y - h_cab, W - 0.4, h_cab, AZUL, AZUL)
    ax.text(W / 2, y - 0.3, "CAPÍTULO III", ha="center", va="center", color="#DDE6F0", fontsize=9)
    ax.text(W / 2, y - 0.68, "Evaluación de Modelos Base", ha="center", va="center", color="white",
            fontsize=14, fontweight="bold")
    ax.text(W / 2, y - 1.2, "Identificar los modelos de Machine Learning y Deep Learning con mejor desempeño y medir\n"
            "cómo se ve afectado ese desempeño por la variabilidad visual entre conjuntos de datos.",
            ha="center", va="center", color="white", fontsize=9, linespacing=1.4)
    y = baja(y - h_cab)

    x_izq, w_izq = 0.2, 2.5
    x_der = x_izq + w_izq
    for (etq, nombre, color, fondo, secciones), h in zip(ETAPAS, alto_etapas):
        caja(x_izq, y - h, W - 0.4, h, fondo, color, lw=2)
        ax.add_patch(plt.Rectangle((x_izq + 0.02, y - h + 0.02), w_izq, h - 0.04, color=color, lw=0))
        ax.text(x_izq + w_izq / 2, y - h / 2 + 0.25 + 0.17 * nombre.count("\n"), etq, ha="center",
                va="center", color="#E8EEF5", fontsize=9)
        ax.text(x_izq + w_izq / 2, y - h / 2 - 0.1, nombre, ha="center", va="center", color="white",
                fontsize=10, fontweight="bold", linespacing=1.25)
        yt = y - pad
        for s in secciones:
            num, titulo, busca, res = s
            ht = alto_tarjeta(s)
            caja(x_der + 0.2, yt - ht, W - x_der - 0.6, ht, "white", color, lw=1.1, r=0.1)
            caja(x_der + 0.4, yt - 0.55, 0.85, 0.4, color, color, lw=0, r=0.08)
            ax.text(x_der + 0.825, yt - 0.35, num, ha="center", va="center", color="white", fontsize=10,
                    fontweight="bold")
            xt = x_der + 1.45
            ax.text(xt, yt - 0.35, titulo, va="center", color=color, fontsize=9.8, fontweight="bold")
            yl = yt - 0.8
            ax.text(xt, yl, "¿Qué se busca?", va="center", color=color, fontsize=8.8, fontweight="bold")
            for ln in lineas(busca):
                yl -= LH
                ax.text(xt, yl, ln, va="center", color="#222222", fontsize=8.8)
            yl -= LH
            ax.text(xt, yl, "Resultado:", va="center", color=color, fontsize=8.8, fontweight="bold")
            for ln in lineas(res):
                yl -= LH
                ax.text(xt, yl, ln, va="center", color="#222222", fontsize=8.8)
            yt -= ht + gap_t
        y = baja(y - h)

    caja(0.2, y - h_pie, W - 0.4, h_pie, "#EEF2F7", AZUL, lw=2)
    ax.text(W / 2, y - 0.35, "Línea base experimental y diagnóstico del efecto de la variabilidad visual",
            ha="center", va="center", color=AZUL, fontsize=10.5, fontweight="bold")
    ax.text(W / 2, y - 0.72, "Insumo para el desarrollo del modelo propuesto (Capítulo IV) y para su "
            "validación (Capítulo V)", ha="center", va="center", color="#222222", fontsize=9.5)

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
    fig.savefig(FIG_OUT, dpi=200, bbox_inches="tight")  # misma figura, usada en Documentacion_Tecnica
    plt.close(fig)
    buf.seek(0)
    return buf


# --- helpers de formato -------------------------------------------------------

def parrafo(doc, texto, bold_prefix=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(texto)
    return p


def vinetas(doc, items):
    for etiqueta, texto in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(f"{etiqueta}: ").bold = True
        p.add_run(texto)


def resultado(doc, texto):
    parrafo(doc, texto, bold_prefix="Resultado de la sección. ")


def rotulo(doc, numero, titulo):
    """Rotulo de tabla/figura en formato APA: numero en negrita y titulo en cursiva."""
    doc.add_paragraph().add_run(numero).bold = True
    doc.add_paragraph().add_run(titulo).italic = True


def fuente(doc, texto="Fuente: Elaboración propia."):
    doc.add_paragraph().add_run(texto).font.size = Pt(9)


def _sombrear(celda, color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    celda._tc.get_or_add_tcPr().append(shd)


def tabla(doc, encabezado, filas, anchos_cm, resaltar=()):
    t = doc.add_table(rows=1, cols=len(encabezado))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(encabezado):
        c = t.rows[0].cells[i]
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size = True, Pt(9.5)
        _sombrear(c, "D9E2F3")
    for k, fila in enumerate(filas):
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            r = celdas[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(9.5)
            if k in resaltar:
                r.bold = True
                _sombrear(celdas[i], "FFF2CC")
    t.autofit = False
    for i, w in enumerate(anchos_cm):  # LibreOffice/Word leen el ancho de tblGrid, no solo de cada celda
        t.columns[i].width = Cm(w)
        for c in t.columns[i].cells:
            c.width = Cm(w)
    return t


# --- contenido ----------------------------------------------------------------

def main():
    intra = cargar_intra()
    por_modelo = {r["model"]: r for r in intra}
    vit, dn = por_modelo["vit_base_patch16_224"], por_modelo["densenet121"]
    caidas = [float(por_modelo[m]["f1"]) - CROSS_DATASET[(m, d)]["f1"]
              for m in ELEGIDOS for d in ("Trucha", "Rohu")]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(2.5))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("Capítulo III: Evaluación de Modelos Base", level=1)

    # 3.1 ---------------------------------------------------------------------
    doc.add_heading("3.1. Introducción", level=2)
    parrafo(doc,
        "El presente capítulo tiene como propósito identificar los modelos de Machine Learning con mejor "
        "desempeño en la clasificación de la frescura del pescado a partir de imágenes oculares y "
        "determinar en qué medida dicho desempeño se ve afectado por la variabilidad visual existente "
        "entre distintos conjuntos de datos. De esta manera, el capítulo da respuesta al primer y segundo "
        "objetivo específico de la investigación y establece la línea base experimental frente a la cual "
        "se comparará el modelo propuesto.")
    parrafo(doc,
        "En relación con el método de investigación (sección 2.4.2), el capítulo desarrolla sus dos "
        "primeras fases. La Fase 1 (línea base), asociada al Objetivo Específico 1, se ejecuta en las "
        "secciones 3.2 a 3.6 y concluye con la selección de los modelos de referencia a partir de su "
        "desempeño intra-dataset. La Fase 2 (diagnóstico), asociada al Objetivo Específico 2, se ejecuta "
        "en las secciones 3.7 y 3.8 y trabaja únicamente con los modelos seleccionados. Finalmente, la "
        "sección 3.9 consolida los productos de ambas fases. Cabe precisar que en las secciones 3.2 y 3.3 "
        "se preparan también los conjuntos de datos de los dominios objetivo y se define el esquema "
        "cross-dataset, insumos que se emplean posteriormente en la Fase 2.")
    parrafo(doc,
        "El capítulo se organiza en cinco etapas secuenciales. En la primera etapa (preparación) se "
        "describen los conjuntos de datos y su procesamiento común, que comprende la depuración, el "
        "recorte de la región ocular, la homologación de las categorías de frescura entre dominios y la "
        "normalización; en esta etapa no se entrena ningún modelo, y se diferencia el conjunto fuente, que "
        "se particiona para el entrenamiento, de los conjuntos objetivo, que se reservan hasta la "
        "evaluación cross-dataset. Asimismo, se define el diseño experimental común que se aplicará a "
        "todos los modelos candidatos. En la segunda etapa (implementación y evaluación intra-dataset) se "
        "entrenan los modelos candidatos en el dominio fuente y se compara su desempeño. En la tercera "
        "etapa (selección) se eligen los modelos de referencia que conforman la línea base experimental. "
        "En la cuarta etapa (evaluación cross-dataset y análisis) se evalúan los modelos de referencia "
        "sobre los dominios objetivo y se cuantifica el efecto de la variabilidad visual sobre su "
        "desempeño y su robustez. Finalmente, en la quinta etapa (consolidación) se sintetizan los "
        "resultados en la línea base experimental, que constituye el insumo para el desarrollo del modelo "
        "propuesto en el Capítulo IV y para su validación en el Capítulo V.")
    parrafo(doc,
        "En la Figura 3.1 se presenta el flujo del capítulo, en el que se resume para cada sección qué se "
        "busca y el resultado esperado.")
    rotulo(doc, "Figura 3.1", "Flujo del Capítulo III: Evaluación de Modelos Base")
    doc.add_picture(figura_flujo(), width=Cm(15.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    fuente(doc)
    parrafo(doc,
        "En la Tabla 3.1 se muestra la correspondencia entre las secciones del capítulo y las fases y "
        "actividades del método de investigación.")
    rotulo(doc, "Tabla 3.1", "Correspondencia entre las secciones del Capítulo III y el método de investigación")
    tabla(doc, ["Sección", "Fase del método", "Actividad del método de investigación que se desarrolla"], [
        ("3.2", "Fase 1 y Fase 2", "Seleccionar y preparar el conjunto de datos base para el entrenamiento "
         "de los modelos (Fase 1). Seleccionar los conjuntos de datos experimentales con diferencias "
         "visuales representativas en iluminación, fondo y condiciones de adquisición (Fase 2)."),
        ("3.3", "Fase 1 y Fase 2", "Identificar y seleccionar los algoritmos con mejor respaldo en la "
         "literatura científica para su evaluación comparativa (Fase 1). Diseñar el esquema experimental "
         "de transferencia entre conjuntos de datos (cross-dataset) (Fase 2)."),
        ("3.4", "Fase 1", "Entrenar los modelos seleccionados bajo condiciones homogéneas de evaluación."),
        ("3.5", "Fase 1", "Comparar el desempeño de los modelos entrenados mediante métricas de clasificación."),
        ("3.6", "Fase 1", "Seleccionar los modelos con mejor desempeño para establecer la línea base "
         "experimental de las fases posteriores."),
        ("3.7", "Fase 2", "Ejecutar los modelos de referencia identificados bajo las dos configuraciones "
         "cross-dataset resultantes."),
        ("3.8", "Fase 2", "Medir la degradación del desempeño mediante métricas de clasificación y métricas "
         "de robustez asociadas a la generalización. Analizar el impacto de la variabilidad visual sobre "
         "el comportamiento de los modelos seleccionados."),
        ("3.9", "Fase 1 y Fase 2", "No incorpora una actividad nueva: consolida los productos de ambas "
         "fases, es decir, la línea base experimental (Fase 1) y el diagnóstico del efecto de la "
         "variabilidad visual (Fase 2)."),
    ], [1.6, 2.8, 11.6])
    fuente(doc, "Nota. La actividad de revisión de la literatura científica de la Fase 1 se desarrolla en "
                "el Capítulo II (Marco Referencial). Fuente: Elaboración propia.")

    # 3.2 ---------------------------------------------------------------------
    doc.add_heading("3.2. Conjuntos de datos y preparación de datos", level=2)
    parrafo(doc,
        "La investigación trabaja con tres repositorios de imágenes oculares de pescado. Mendeley Data "
        "(DataSetNotClean, DNC) es el dominio fuente: con él se entrenan y comparan los modelos "
        "candidatos. OneDrive (trucha, Perú) y Kaggle (rohu, India) son los dominios objetivo: provienen "
        "de otras especies, cámaras e iluminación, y se reservan para medir la generalización de los "
        "modelos de referencia.")
    vinetas(doc, [
        ("Mendeley (fuente)", "8,000 imágenes balanceadas, 4,000 fresco y 4,000 no_fresco, de 8 especies, "
         "ya recortadas al ojo (aprox. 300×300 a 440×440 px)."),
        ("OneDrive – Trucha (objetivo 1)", "1,150 imágenes tras la depuración; fotos originales de cabeza "
         "completa a resolución de cámara (aprox. 4160×3120 px)."),
        ("Kaggle – Rohu (objetivo 2)", "1,261 imágenes tras la depuración; fotos originales de cabeza "
         "completa (aprox. 4032×3024 px)."),
    ])
    parrafo(doc,
        "Los tres repositorios reciben un procesamiento común de cuatro pasos, para que las diferencias "
        "entre dominios se deban a las condiciones visuales y no al encuadre de las imágenes:")
    vinetas(doc, [
        ("Depuración", "en OneDrive y Kaggle se conservan las detecciones del ojo con estado «ok» y las "
         "rescatadas en revisión manual; en Kaggle se excluyen además los falsos positivos confirmados "
         "(escama, aleta o imagen borrosa)."),
        ("Recorte de la región ocular", "el ojo se detecta con Fast Radial Symmetry Transform y se recorta "
         "un cuadrado centrado en él, con el mismo encuadre que Mendeley."),
        ("Homologación de categorías", "todas las etiquetas se llevan a un problema binario: fresco "
         "(Mendeley: «Highly Fresh» y «Fresh»; OneDrive: días 1 y 3; Kaggle: Fresh_Eyes) y no_fresco "
         "(Mendeley: «Not Fresh»; OneDrive: días 5, 7 y 9; Kaggle: Nonfresh_Eyes)."),
        ("Normalización", "brillo y contraste con CLAHE en los dominios objetivo, redimensionado a 330×330 "
         "px y, al ingresar al modelo, normalización con la media y desviación de ImageNet."),
    ])
    parrafo(doc,
        "En esta sección no se entrena ningún modelo. Mendeley se particiona para el entrenamiento "
        "(sección 3.3) y solo su conjunto de entrenamiento recibe data augmentation. OneDrive y Kaggle no "
        "intervienen en el entrenamiento ni en la selección de los modelos base: se reservan hasta la "
        "evaluación cross-dataset (sección 3.7).")
    resultado(doc, "conjuntos de datos comparables entre dominios, sin entrenamiento previo.")

    # 3.3 ---------------------------------------------------------------------
    doc.add_heading("3.3. Diseño experimental de los modelos candidatos", level=2)
    parrafo(doc,
        "A partir de la revisión de la literatura (Capítulo II) se seleccionan diez modelos candidatos que "
        "cubren las dos familias de arquitecturas usadas en clasificación de imágenes: Transformer "
        "(ViT-Base/16, Swin-Base y DINOv2-Small) y redes convolucionales, CNN (DenseNet-121, "
        "EfficientNet-B0, EfficientNet-Lite0, MobileNetV3-Small, ResNet-50, MobileNetV2 y MobileNetV1). "
        "Así, la selección posterior no depende de una sola familia.")
    doc.add_heading("Partición de los datos", level=3)
    parrafo(doc,
        "Las 8,000 imágenes de Mendeley se dividen con StratifiedKFold (n_splits=5, shuffle=True, "
        "random_state=42) en cinco bloques de 1,600 imágenes. En cada una de las cinco corridas, un bloque "
        "se reserva como prueba (20%) y del 80% restante se separa un 15% para validación con "
        "train_test_split estratificado y la misma semilla. La distribución efectiva es 68% entrenamiento "
        "(5,440 imágenes), 12% validación (960) y 20% prueba (1,600). Al terminar las cinco corridas, cada "
        "imagen fue prueba exactamente una vez.")
    vinetas(doc, [
        ("Entrenamiento (68%)", "actualiza los pesos del modelo."),
        ("Validación (12%)", "calcula val_F1 en cada epoch y elige el mejor checkpoint; no actualiza pesos."),
        ("Prueba (20%)", "solo mide el checkpoint ya elegido; nunca se usa para tomar decisiones."),
    ])
    parrafo(doc,
        "Limitación: la partición se hace a nivel de imagen individual, sin agrupar imágenes de una misma "
        "captura o ejemplar. Si el conjunto contiene varias imágenes de una misma captura, pueden quedar en "
        "conjuntos distintos y elevar las métricas de prueba.")
    doc.add_heading("Hiperparámetros", level=3)
    vinetas(doc, [
        ("Inicialización", "pesos preentrenados (ImageNet; en DINOv2-Small, pesos auto-supervisados de "
         "DINOv2) y fine-tuning completo de toda la arquitectura."),
        ("Optimización", "AdamW (learning rate 1e-4, weight decay 0.05); 3 epochs de warmup lineal y luego "
         "scheduler coseno; en los Transformer, el backbone usa un learning rate 10 veces menor que la cabeza."),
        ("Pérdida", "Focal Loss (γ = 2) con label smoothing 0.1; gradient clipping con norma máxima 1.0."),
        ("Entrada", "CNN: 256×256 px, batch size 16. Transformer: 224×224 px, batch size 8."),
        ("Duración", "máximo 30 epochs con early stopping de paciencia 7."),
        ("Data augmentation (solo entrenamiento)", "recorte aleatorio, volteos, variación de color, escala "
         "de grises, desenfoque, rotación de ±15° y borrado aleatorio."),
        ("Semilla", "42 en todas las particiones y entrenamientos. Algunas operaciones en MPS (Apple M2) no "
         "son deterministas, por lo que un reentrenamiento puede variar en décimas."),
    ])
    doc.add_heading("Métricas", level=3)
    parrafo(doc,
        "A partir de la matriz de confusión se calculan exactitud (Accuracy), precisión, sensibilidad "
        "(Recall), especificidad y F1-score. Las métricas por clase se calculan tomando cada clase como "
        "positiva por turno y se promedian (macro), de modo que fresco y no_fresco pesan igual. La métrica "
        "central es el F1 macro. Con las probabilidades de salida se construye la curva ROC y se calcula "
        "el AUC. La estabilidad entre particiones se mide con la desviación estándar del F1 de los cinco "
        "folds.")
    doc.add_heading("Escenarios", level=3)
    vinetas(doc, [
        ("Intra-dataset", "entrenamiento y prueba en Mendeley (secciones 3.4 a 3.6), con los diez candidatos."),
        ("Cross-dataset", "Mendeley → OneDrive (Trucha) y Mendeley → Kaggle (Rohu), solo con los modelos de "
         "referencia (secciones 3.7 y 3.8)."),
    ])
    resultado(doc, "protocolo experimental común para los diez candidatos.")

    # 3.4 ---------------------------------------------------------------------
    doc.add_heading("3.4. Implementación de los modelos candidatos", level=2)
    parrafo(doc,
        "Cada uno de los diez modelos se entrena en cada uno de los cinco folds: 50 entrenamientos "
        "modelo–fold, todos bajo el protocolo de la sección 3.3. La implementación usa Python con PyTorch y "
        "torchvision, timm como repositorio de arquitecturas, scikit-learn para la partición y las "
        "métricas, y pandas, numpy y matplotlib para el tratamiento de resultados. El entorno de ejecución "
        "es una MacBook con chip Apple M2 y aceleración Metal Performance Shaders (MPS).")
    parrafo(doc,
        "En cada epoch se calcula val_F1. Si supera el mejor valor registrado en ese fold, el checkpoint se "
        "guarda como el nuevo mejor modelo. Si se acumulan siete epochs sin mejora, el entrenamiento se "
        "detiene antes de los 30 epochs (early stopping). Al terminar, se recarga el mejor checkpoint, que "
        "no es necesariamente el del último epoch.")
    parrafo(doc,
        "Los pesos se aprenden con entrenamiento, el checkpoint se elige con validación y el desempeño se "
        "mide con prueba. Esta separación evita usar la prueba para tomar decisiones durante el "
        "entrenamiento.")
    parrafo(doc,
        "Cada corrida experimental corresponde a un fold: el modelo se entrena con cuatro folds "
        "(entrenamiento + validación) y se evalúa una vez sobre el fold restante (prueba). La Tabla 3.2 "
        "resume las corridas de cada modelo.")
    rotulo(doc, "Tabla 3.2", "Corridas experimentales por modelo candidato")
    ep = epochs_por_modelo()
    filas = []
    for r in intra:
        folds = ep.get(r["model"], {})
        if len(folds) == 5:
            txt = str(sum(folds.values()))
        elif folds:
            txt = f"{sum(folds.values())} documentados (folds {min(folds)}–{max(folds)})"
        else:
            txt = "sin registro en el log"
        filas.append((NOMBRES[r["model"]], "5", txt, "5", "1,600 × 5 = 8,000"))
    filas.append(("Total", "50", "—", "50", "80,000"))
    tabla(doc, ["Modelo", "Corridas (folds)", "Epochs ejecutados", "Evaluaciones de prueba",
                "Predicciones OOF"], filas, [3.5, 2.2, 4.4, 2.6, 3.3], resaltar=(len(filas) - 1,))
    fuente(doc, "Nota. Epochs ejecutados = suma de los 5 folds (máximo 30 por fold; menos de 150 indica "
                "early stopping: Swin-Base se detuvo en el epoch 25 del fold 4 y ResNet-50 en el 28 del fold "
                "3). El entrenamiento se interrumpió y se reanudó; el log conservado "
                f"({LOG_ENTRENAMIENTO}) no incluye ViT-Base/16 folds 1 y 2 ni EfficientNet-Lite0 y "
                "MobileNetV3-Small. Sus checkpoints y resultados de prueba sí existen y forman parte de las "
                "Tablas 3.3 y siguientes; solo falta la curva de val_F1 por epoch.")
    vinetas(doc, [
        ("Entrenamientos", "10 modelos × 5 folds = 50."),
        ("Checkpoint que pasa a prueba", "el de mejor val_F1 macro del fold."),
        ("Detalle por modelo", "curvas de val_F1 por epoch y resultados por fold en "
         "Documentacion_Tecnica_10_Modelos.docx, sección 4."),
    ])
    resultado(doc, "50 checkpoints entrenados, cinco por modelo candidato.")

    # 3.5 ---------------------------------------------------------------------
    doc.add_heading("3.5. Evaluación intra-dataset de los modelos candidatos", level=2)
    parrafo(doc,
        "Cada checkpoint se evalúa una sola vez sobre la prueba de su fold. Como los cinco conjuntos de "
        "prueba son disjuntos y cubren las 8,000 imágenes, sus predicciones se concatenan en 8,000 "
        "predicciones out-of-fold (OOF): cada imagen se predice con un modelo que no la vio en el "
        "entrenamiento. Sobre ellas se recalculan las métricas globales (pooled); la matriz de confusión "
        "pooled es la suma de las cinco matrices de prueba. Además, se reporta el F1 de los cinco folds "
        "como media ± desviación estándar.")
    rotulo(doc, "Tabla 3.3", "Desempeño intra-dataset de los diez modelos candidatos en Mendeley (pooled, "
                             "8,000 predicciones OOF)")
    filas, resaltar = [], []
    for k, r in enumerate(intra):
        if r["model"] in ELEGIDOS:
            resaltar.append(k)
        filas.append((NOMBRES[r["model"]], *(f'{float(r[c]):.2f}' for c in
                                             ("accuracy", "precision", "recall", "f1", "specificity")),
                      f'{float(r["auc"]):.4f}', f'{r["f1_media"]:.2f} ± {r["f1_de"]:.2f}'))
    tabla(doc, ["Modelo", "Exactitud (%)", "Precisión (%)", "Sensibilidad (%)", "F1 macro (%)",
                "Especificidad (%)", "AUC", "F1 por fold (%)"], filas,
          [3.3, 1.7, 1.7, 2.0, 1.6, 2.1, 1.5, 2.1], resaltar=resaltar)
    fuente(doc, "Nota. Ordenado por F1 macro. Resaltados los modelos seleccionados en la sección 3.6. "
                f"Fuente: {CSV_POOLED} y {CSV_FOLDS}.")
    parrafo(doc,
        f"ViT-Base/16 obtiene el mayor F1 macro ({vit['f1']}%) y DenseNet-121 el segundo ({dn['f1']}%), "
        f"con el mayor AUC de los diez ({float(dn['auc']):.4f}). Los cuatro mejores modelos quedan entre "
        "95.4% y 97.0% de F1; los seis restantes, entre 89.0% y 94.3%. Como el conjunto está balanceado, "
        "exactitud, sensibilidad y especificidad coinciden prácticamente con el F1.")
    resultado(doc, "desempeño comparativo de los diez modelos candidatos. Comparación completa en "
                   "Documentacion_Tecnica_10_Modelos.docx, sección 3; matrices de confusión, curvas ROC y "
                   "cálculos por fold de cada modelo en su sección 4.")

    # 3.6 ---------------------------------------------------------------------
    doc.add_heading("3.6. Selección de los modelos de referencia", level=2)
    parrafo(doc,
        "La selección no depende de una sola métrica, sino de cinco criterios complementarios:")
    vinetas(doc, [
        ("Desempeño", f"F1 macro pooled: ViT-Base/16 es primero ({vit['f1']}%) y DenseNet-121 segundo "
         f"({dn['f1']}%); Swin-Base ({por_modelo['swin_base_patch4_window7_224']['f1']}%) y DINOv2-Small "
         f"({por_modelo['vit_small_patch14_dinov2']['f1']}%) quedan después."),
        ("Discriminación", f"DenseNet-121 tiene el mayor AUC de los diez ({float(dn['auc']):.4f}) y "
         f"ViT-Base/16 el segundo ({float(vit['auc']):.4f})."),
        ("Estabilidad", f"desviación estándar del F1 entre folds: ViT-Base/16 {vit['f1_de']:.2f}, la más "
         f"baja entre los modelos de alto desempeño; DenseNet-121 {dn['f1_de']:.2f}."),
        ("Eficiencia", f"DenseNet-121 (~{PARAMS_M['densenet121']}M parámetros) logra más del 98% del F1 de "
         f"ViT-Base/16 (~{PARAMS_M['vit_base_patch16_224']}M) con alrededor del 8% de sus parámetros."),
        ("Diversidad arquitectónica", "se busca un representante Transformer y uno CNN. Swin-Base tiene un "
         "F1 cercano a DenseNet-121, pero es Transformer como ViT-Base/16 y no aporta la comparación ViT "
         "frente a CNN."),
    ])
    parrafo(doc,
        "Con estos criterios se seleccionan ViT-Base/16 (Transformer) y DenseNet-121 (CNN) como modelos de "
        "referencia. Ninguno de los dos es el modelo propuesto de la tesis: ambos forman la línea base "
        "experimental frente a la cual se comparará el modelo propuesto.")
    resultado(doc, "modelos de referencia ViT-Base/16 y DenseNet-121 (línea base experimental). Sustento "
                   "detallado en Sustento_Seleccion_Modelos.docx.")

    # 3.7 ---------------------------------------------------------------------
    doc.add_heading("3.7. Evaluación cross-dataset de los modelos de referencia", level=2)
    parrafo(doc,
        "Los dos modelos de referencia, entrenados con Mendeley, se llevan a los dominios objetivo con el "
        "método congelado (linear probe). El backbone conserva los pesos aprendidos en Mendeley y no se "
        "actualiza: sus parámetros no reciben gradiente y permanece en modo de evaluación, de modo que "
        "BatchNorm no cambia sus estadísticas. Solo se reentrena la cabeza de clasificación con imágenes "
        "del dominio objetivo. Así, el resultado mide cuán transferibles son las características "
        "aprendidas en Mendeley.")
    vinetas(doc, [
        ("Escenarios", "Mendeley → OneDrive (Trucha, 1,150 imágenes) y Mendeley → Kaggle (Rohu, 1,261 "
         "imágenes)."),
        ("Partición", "5 folds con semilla 42: StratifiedKFold en Trucha y StratifiedGroupKFold en Rohu, "
         "agrupando fotos casi idénticas para que no queden a la vez en entrenamiento y validación."),
        ("Punto de partida", "en el fold k se carga el checkpoint de Mendeley del mismo fold k."),
        ("Entrenamiento de la cabeza", "AdamW (learning rate 1e-4, weight decay 1e-4), scheduler coseno, "
         "entropía cruzada, máximo 30 epochs con early stopping de paciencia 8; data augmentation sin "
         "duplicar el conjunto de entrenamiento (sin x2)."),
        ("Evaluación", "sobre la validación del fold; el F1 oficial es la media de los cinco folds."),
    ])
    rotulo(doc, "Tabla 3.4", "Desempeño cross-dataset de los modelos de referencia (método congelado)")
    filas = []
    for m in ELEGIDOS:
        for dom in ("Trucha", "Rohu"):
            c = CROSS_DATASET[(m, dom)]
            filas.append((NOMBRES[m], f"Mendeley → {dom}", f"{c['f1']:.2f}", f"{c['f1_pool']:.2f}",
                          f"{c['acc_pool']:.2f}", f"{c['auc_pool']:.4f}"))
    tabla(doc, ["Modelo", "Escenario", "F1 macro (%, media 5 folds)", "F1 pooled (%)", "Exactitud pooled (%)",
                "AUC pooled"], filas, [3.0, 3.2, 2.8, 2.2, 2.6, 2.0])
    fuente(doc, "Nota. Los valores pooled provienen de una corrida aparte con la misma configuración, por eso "
                "difieren levemente del F1 oficial. Fuente: repositorio Generalizacion_ViT_DenseNet "
                "(3_resultados/ y Documentacion_Tecnica_ViT_DenseNet.docx).")
    resultado(doc, "desempeño de los modelos de referencia frente a dominios no vistos durante el "
                   "entrenamiento base.")

    # 3.8 ---------------------------------------------------------------------
    doc.add_heading("3.8. Análisis del efecto de la variabilidad visual entre conjuntos de datos", level=2)
    parrafo(doc,
        "El efecto del cambio de dominio se cuantifica comparando el F1 macro de cada modelo en Mendeley "
        "(intra-dataset) con su F1 macro en cada dominio objetivo (cross-dataset), mediante tres "
        "indicadores de robustez:")
    vinetas(doc, [
        ("Tasa de Retención del Desempeño", "TRD = (F1_objetivo / F1_fuente) × 100. Porcentaje del "
         "desempeño de Mendeley que se conserva; 100% significa sin caída."),
        ("Brecha de error", "Gap = Error_objetivo − Error_fuente, con Error = 1 − F1/100. Cuánto aumenta "
         "el error al cambiar de dominio."),
        ("Reducción de la Brecha de Error", "RBE = (Gap_DenseNet − Gap_modelo) / Gap_DenseNet × 100. "
         "Cuánto reduce un modelo la brecha frente a DenseNet-121 (baseline CNN); RBE > 0 indica mejor "
         "transferencia."),
    ])
    rotulo(doc, "Tabla 3.5", "Degradación y robustez de los modelos de referencia ante el cambio de dominio")
    filas = []
    for m in ELEGIDOS:
        f1_src = float(por_modelo[m]["f1"])
        for dom in ("Trucha", "Rohu"):
            c = CROSS_DATASET[(m, dom)]
            filas.append((NOMBRES[m], dom, f"{f1_src:.2f}", f"{c['f1']:.2f}", f"{c['f1'] - f1_src:+.2f}",
                          f"{c['trd']:.2f}", f"{c['gap']:.4f}", "—" if c["rbe"] is None else f"{c['rbe']:.2f}"))
    tabla(doc, ["Modelo", "Dominio objetivo", "F1 Mendeley (%)", "F1 objetivo (%)", "Variación (pp)",
                "TRD (%)", "Gap", "RBE (%)"], filas, [2.8, 2.0, 2.0, 2.0, 1.9, 1.7, 1.6, 1.6])
    fuente(doc, "Nota. RBE se calcula respecto a DenseNet-121, por lo que no aplica al propio baseline. "
                "Fuente: Indicadores_Robustez_TRD_Gap_RBE.docx (repositorio Generalizacion_ViT_DenseNet).")
    vt, vr = CROSS_DATASET[("vit_base_patch16_224", "Trucha")], CROSS_DATASET[("vit_base_patch16_224", "Rohu")]
    dt, dr = CROSS_DATASET[("densenet121", "Trucha")], CROSS_DATASET[("densenet121", "Rohu")]
    parrafo(doc,
        f"Ambos modelos pierden desempeño al cambiar de dominio. ViT-Base/16 cae de {vit['f1']}% a "
        f"{vt['f1']:.2f}% en Trucha y a {vr['f1']:.2f}% en Rohu; DenseNet-121 cae de {dn['f1']}% a "
        f"{dt['f1']:.2f}% y {dr['f1']:.2f}%. Como el procesamiento de la sección 3.2 iguala el encuadre y "
        "la resolución, esa pérdida se asocia a la variabilidad visual entre conjuntos (iluminación, "
        "cámara, fondo y especie).")
    parrafo(doc,
        f"ViT-Base/16 es más robusto que DenseNet-121 en los dos dominios: retiene más desempeño (TRD "
        f"{vt['trd']:.2f}% frente a {dt['trd']:.2f}% en Trucha y {vr['trd']:.2f}% frente a {dr['trd']:.2f}% "
        f"en Rohu), tiene una brecha de error menor y su RBE es positivo ({vt['rbe']:.2f}% en Trucha y "
        f"{vr['rbe']:.2f}% en Rohu). Para los dos modelos, la caída es mayor en Trucha que en Rohu.")
    parrafo(doc,
        "El cambio de dominio incluye también una diferencia de especie (Mendeley reúne 8 especies de "
        "Indonesia; OneDrive es trucha de Perú y Kaggle es rohu de India). Por ello, la degradación medida "
        "combina variabilidad de adquisición y diferencias propias de cada especie.")
    resultado(doc, "magnitud del efecto del cambio de dominio y robustez comparada de los modelos de "
                   "referencia.")

    # 3.9 ---------------------------------------------------------------------
    doc.add_heading("3.9. Consolidación de la línea base experimental", level=2)
    parrafo(doc,
        "La Tabla 3.6 reúne los resultados de ambas fases para los dos modelos de referencia. Constituye "
        "la línea base experimental frente a la cual se comparará el modelo propuesto en el Capítulo IV y "
        "se validará en el Capítulo V.")
    rotulo(doc, "Tabla 3.6", "Línea base experimental: resultados intra-dataset, cross-dataset y de retención")
    filas = []
    for m in ELEGIDOS:
        r = por_modelo[m]
        t, ro = CROSS_DATASET[(m, "Trucha")], CROSS_DATASET[(m, "Rohu")]
        filas.append((NOMBRES[m], r["f1"], f'{float(r["auc"]):.4f}', f"{t['f1']:.2f}", f"{ro['f1']:.2f}",
                      f"{t['trd']:.2f}", f"{ro['trd']:.2f}"))
    tabla(doc, ["Modelo", "F1 Mendeley (%)", "AUC Mendeley", "F1 Trucha (%)", "F1 Rohu (%)",
                "TRD Trucha (%)", "TRD Rohu (%)"], filas, [3.0, 2.3, 2.0, 2.2, 2.2, 2.2, 2.1])
    fuente(doc)
    vinetas(doc, [
        ("Producto de la Fase 1", "ViT-Base/16 y DenseNet-121 como modelos de referencia, con F1 macro de "
         f"{vit['f1']}% y {dn['f1']}% en el dominio fuente."),
        ("Producto de la Fase 2", f"diagnóstico del problema: los modelos pierden entre {min(caidas):.2f} y "
         f"{max(caidas):.2f} puntos de F1 al cambiar de dominio, y ViT-Base/16 transfiere mejor que "
         "DenseNet-121."),
        ("Uso posterior", "el modelo propuesto deberá superar esta línea base, en especial en los "
         "escenarios cross-dataset."),
    ])
    resultado(doc, "línea base experimental consolidada y diagnóstico del efecto de la variabilidad visual.")

    doc.save(OUT)
    print(f"Guardado en {OUT}")


if __name__ == "__main__":
    main()
